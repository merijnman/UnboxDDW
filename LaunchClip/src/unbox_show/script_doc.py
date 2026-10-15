"""Readable script of the show as a Word file (`unbox script`), for the team to read and edit.

No timings or file names: per section who says what, with the delivery directions
([excited] and so on) in italics, the Dutch subtitle below each line, and short
grey notes for what is on screen and what happens in the room. Changes the team
makes in Word are copied back into show.yaml by hand.
"""

from __future__ import annotations

import re
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from .model import Entry, ShowFile, format_time

SECTION_RE = re.compile(r"^\s*#\s*-{2,}\s*(\d+\.\s.*?)\s*-{2,}")
TAG_RE = re.compile(r"(\[[^\]]*\])")
QUIET_SOUNDS = {"radio_crackle", "timer_end"}  # not worth a note in a script

PURPLE = RGBColor(0x45, 0x37, 0x58)
GREY = RGBColor(0x80, 0x80, 0x80)
NL_COLOUR = RGBColor(0x8A, 0x6D, 0x1E)


def sections(path: Path) -> dict[int, str]:
    """Timeline index -> title of the section that starts there, from the `# ---- 1. Title ---` comments."""
    text = path.read_text(encoding="utf-8")
    titles = [(no, m.group(1)) for no, line in enumerate(text.splitlines()) if (m := SECTION_RE.match(line))]
    root = yaml.compose(text)
    timeline = next(value for key, value in root.value if key.value == "timeline")
    result: dict[int, str] = {}
    for index, node in enumerate(timeline.value):
        line = node.start_mark.line
        before = [title for no, title in titles if no < line]
        if before and before[-1] not in result.values():
            result[index] = before[-1]
    return result


def _note(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.italic, run.font.size, run.font.color.rgb = True, Pt(9), GREY


def _spoken(paragraph, text: str) -> None:
    """Line text with [audio tags] as small italic directions."""
    for part in TAG_RE.split(text):
        if not part.strip():
            if part:
                paragraph.add_run(part)
            continue
        if TAG_RE.fullmatch(part):
            run = paragraph.add_run(f"({part[1:-1]})")
            run.italic, run.font.color.rgb = True, GREY
        else:
            paragraph.add_run(part)


def _notes_for(show: ShowFile, item) -> list[str]:
    notes = []
    screen = show.screens.get(item.screen or "")
    if screen and screen.text_en:
        notes.append(f"Screen: {screen.text_en}".replace("\n", " / "))
    if item.sound and item.sound not in QUIET_SOUNDS:
        notes.append(f"Sound: {item.sound.replace('_', ' ')}")
    if item.ops:
        notes.append(f"Room: {item.ops}")
    return notes


def _entry(doc, show: ShowFile, entry: Entry) -> None:
    for note in _notes_for(show, entry):
        _note(doc, note)
    if entry.timer:
        seconds = entry.timer.seconds or 0
        what = f"{int(seconds // 60)} minutes" if seconds >= 60 else f"{int(seconds)} seconds"
        _note(doc, f"Timer on screen: {what} for the crews" if entry.timer.seconds else "Timer on screen")
    if not entry.say:
        return
    voice = show.voices[entry.say]
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(0)
    label = heading.add_run(voice.label + ("  (muttered aside)" if entry.delivery == "aside" else ""))
    label.bold, label.font.size, label.font.color.rgb = True, Pt(9), PURPLE
    if entry.pace_seconds:  # the countdown: one short line per second, as one paragraph
        _spoken(doc.add_paragraph(), " ".join(line.subtitle or line.en for line in entry.lines))
        for line in entry.lines:
            for note in _notes_for(show, line):
                _note(doc, note)
        return
    for line in entry.lines:
        for note in _notes_for(show, line):
            _note(doc, note)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        _spoken(paragraph, line.subtitle or line.en)
        if line.nl:
            paragraph.add_run().add_break(WD_BREAK.LINE)
            nl = paragraph.add_run(line.nl)
            nl.font.size, nl.font.color.rgb = Pt(9), NL_COLOUR


def write(show: ShowFile, show_path: Path, out: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    doc.add_heading(show.show.title, level=0)
    intro = doc.add_paragraph(
        "The full script of the launch video, made from show.yaml. Edit the text freely, "
        "or add comments: new lines, other wording, more explanation of a task. "
        "Changes are copied back into the video by hand."
    )
    intro.paragraph_format.space_after = Pt(4)
    for text in (
        "(italic in brackets) is how a line is spoken, for the voice: excited, dryly, and so on.",
        "Below each line, in smaller type, is the Dutch subtitle.",
        "Grey notes say what is on screen and what happens in the room (lights, smoke, timers).",
        f"The show lasts about {format_time(show.show.duration_s)} minutes; the crews work at their tables during the timers.",
    ):
        doc.add_paragraph(text, style="List Bullet")
    cast = doc.add_paragraph()
    cast.add_run("Voices: ").bold = True
    cast.add_run(", ".join(voice.label.title() for voice in show.voices.values()) + ".")

    titles = sections(show_path)
    for index, entry in enumerate(show.timeline):
        if index in titles:
            doc.add_heading(re.sub(r"^\d+\.\s*", "", titles[index]), level=1)
        _entry(doc, show, entry)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out
